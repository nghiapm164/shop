from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

slides_data = [
    ("XÂY DỰNG PHẦN MỀM WEBSITE BÁN QUẦN ÁO THỂ THAO NAM", "Sinh viên: Lê Tuấn Hùng, Trần Anh Đăng, Hoàng Văn Tú\nHọc viện Nông nghiệp Việt Nam - Khoa CNTT\nĐơn vị thực tập: Công ty bảo hiểm bưu điện Âu Lạc"),
    ("Nội dung chính", "1. Tổng quan dự án\n2. Công nghệ sử dụng\n3. Phân tích thiết kế hệ thống\n4. Kết quả triển khai\n5. Tổng kết & Hướng phát triển"),
    ("Tổng quan môi trường thực tập", "• Đơn vị: Công ty bảo hiểm bưu điện Âu Lạc\n• Lĩnh vực: Phát triển phần mềm, giải pháp chuyển đổi số\n• Quy trình làm việc: Agile/Scrum (Sprint 2 tuần)\n• Quản lý mã nguồn: Git/GitHub"),
    ("Đặt vấn đề & Mục tiêu", "THỰC TRẠNG:\n• Cửa hàng truyền thống khó quản lý hàng hóa đa dạng biến thể (size, màu).\n• Tốc độ website cũ chậm, UX/UI kém trên thiết bị di động.\n\nMỤC TIÊU:\n• Xây dựng hệ thống TMĐT Mobile-First hiện đại.\n• Quản trị tồn kho và biến thể sản phẩm chặt chẽ.\n• Phân quyền linh hoạt bằng Role-Based Access Control."),
    ("Tổ chức nhóm", "• Lê Tuấn Hùng: Backend (Thiết kế Database, lập trình API).\n• Trần Anh Đăng: Frontend (Thiết kế UI/UX, lập trình HTML/CSS/JS).\n• Hoàng Văn Tú: Quản lý tài liệu, chức năng Auth, Testing & Tối ưu."),
    ("Kiến trúc hệ thống", "• Mô hình: MVC (Model - View - Controller)\n• Kiến trúc thành phần phản hồi (Reactive Component Architecture)\n• Bảo mật hệ thống:\n  - RBAC (Phân quyền theo vai trò)\n  - Chống CSRF (Cross-Site Request Forgery)\n  - Chống SQL Injection (Prepared Statements)"),
    ("Công nghệ và Môi trường áp dụng", "• Backend: Laravel, Livewire\n• Frontend: Tailwind CSS, JS (ES6+)\n• Cơ sở dữ liệu: MySQL\n• Công cụ Admin: Filament PHP (TALL stack)\n• Môi trường & IDE: XAMPP, VS Code, Git"),
    ("Đặc tả yêu cầu hệ thống", "KHÁCH HÀNG:\n• Quản lý tài khoản, xem sản phẩm theo bộ lọc.\n• Giỏ hàng, đặt hàng thanh toán COD.\n• Xem lịch sử và tra cứu đơn hàng.\n\nQUẢN TRỊ VIÊN:\n• Quản lý Sản phẩm (ma trận biến thể Size/Color).\n• Quản lý Danh mục, Đơn hàng, Mã giảm giá (Coupon).\n• Dashboard thống kê doanh thu theo thời gian thực."),
    ("Biểu đồ Use Case tổng quát", "(Bạn hãy chèn hình ảnh Biểu đồ Use Case từ báo cáo Word vào đây)"),
    ("Thiết kế Cơ sở dữ liệu (ERD)", "Hệ thống bao gồm các bảng chính để giải quyết bài toán thương mại điện tử:\n• products, product_variants (Quản lý biến thể màu/size)\n• orders, order_items (Quản lý hóa đơn)\n• users, addresses, flash_sales, coupons...\n\n(Bạn hãy chèn Sơ đồ Thực thể Liên kết - ERD vào đây)"),
    ("Luồng xử lý cốt lõi", "(Bạn hãy chèn Biểu đồ Activity Diagram / Sequence Diagram luồng Đặt hàng vào đây)"),
    ("Giao diện Khách hàng - Mua sắm", "(Chèn ảnh Trang chủ, Danh mục sản phẩm, Giao diện lọc sản phẩm vào đây)"),
    ("Giao diện Khách hàng - Đặt hàng", "(Chèn ảnh Chi tiết sản phẩm -> Giỏ hàng -> Trang Thanh toán vào đây)"),
    ("Giao diện Quản trị - Thống kê", "(Chèn ảnh Dashboard thống kê doanh thu, tổng số đơn hàng vào đây)"),
    ("Giao diện Quản trị - Vận hành", "(Chèn ảnh trang Quản lý sản phẩm, trang xử lý Đơn hàng vào đây)"),
    ("Đánh giá & Kiểm thử (Testing)", "Thực hiện kiểm thử (Test Cases) thành công với các luồng trọng yếu:\n• Chức năng Authentication (Đăng nhập, phân quyền).\n• Giỏ hàng (Kiểm tra giới hạn số lượng tồn kho).\n• Thanh toán và áp dụng mã giảm giá."),
    ("Kết luận & Bài học kinh nghiệm", "KẾT QUẢ:\n• Website load nhanh, giao diện Responsive thân thiện.\n• Quản lý kho hàng chuyên nghiệp, khắc phục nhược điểm hệ thống cũ.\n\nBÀI HỌC KINH NGHIỆM:\n• Kỹ năng quản lý source code bằng Git (chia nhánh, review code).\n• Kinh nghiệm tối ưu hóa Database và làm việc với Laravel ORM."),
    ("Hướng phát triển tương lai", "• Tích hợp cổng thanh toán trực tuyến (VNPAY, MoMo...).\n• Tích hợp hệ thống AI gợi ý sản phẩm theo thói quen mua sắm.\n• Phát triển phiên bản Ứng dụng di động (Mobile App)."),
    ("Lời Cảm Ơn", "Chân thành cảm ơn Thầy Cô, Nhà trường và Công ty Âu Lạc đã tận tình hướng dẫn, hỗ trợ nhóm hoàn thành đề tài!\n\n---- Q&A ----")
]

# Title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = slides_data[0][0]
subtitle.text = slides_data[0][1]

# Content slides
for title_text, content_text in slides_data[1:]:
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    content_placeholder = slide.placeholders[1]
    
    title_placeholder.text = title_text
    content_placeholder.text = content_text

prs.save("Bao_Cao_TTNN.pptx")
